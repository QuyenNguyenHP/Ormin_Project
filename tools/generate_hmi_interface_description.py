from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "HMI" / "public" / "DRUMS_logo.png"
OUTPUT_DIR = ROOT / "Customer_Documents"
OUTPUT = OUTPUT_DIR / "ORMIN_HMI_Interface_Description_Rev0.1.docx"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F2F2"
MID_GREY = "7F8C8D"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edges[edge].items():
            element.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(section, document):
    header = section.header
    header.is_linked_to_previous = False
    table = header.add_table(rows=1, cols=2, width=Cm(17.0))
    table.autofit = False
    table.columns[0].width = Cm(2.2)
    table.columns[1].width = Cm(14.8)
    left, right = table.rows[0].cells
    left.width = Cm(2.2)
    right.width = Cm(14.8)
    if LOGO.exists():
        left.paragraphs[0].add_run().add_picture(str(LOGO), width=Cm(1.15))
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right_p.add_run("HMI INTERFACE DESCRIPTION")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    for cell in (left, right):
        set_cell_border(cell, bottom={"val": "single", "sz": "8", "color": NAVY})

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ORMIN / ORMECO HMI  |  Rev. 0.1  |  Page ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr_text)
    r._r.append(fld_char2)


def add_title(document, text, level=1):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14 if level == 1 else 11)
    r.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), BLUE)
        pbdr.append(bottom)
        p_pr.append(pbdr)
    return p


def add_body(document, text, bold_prefix=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.add_run(item)


def add_screenshot_placeholder(document, screen, description):
    table = document.add_table(rows=2, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17.0)
    title_cell = table.cell(0, 0)
    title_cell.width = Cm(17.0)
    set_cell_shading(title_cell, NAVY)
    set_cell_margins(title_cell, top=80, bottom=80)
    p = title_cell.paragraphs[0]
    r = p.add_run(screen)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(WHITE)

    box = table.cell(1, 0)
    box.width = Cm(17.0)
    table.rows[1].height = Cm(6.25)
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    box.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(box, LIGHT_GREY)
    set_cell_border(
        box,
        top={"val": "dashed", "sz": "8", "color": "A6A6A6"},
        bottom={"val": "dashed", "sz": "8", "color": "A6A6A6"},
        left={"val": "dashed", "sz": "8", "color": "A6A6A6"},
        right={"val": "dashed", "sz": "8", "color": "A6A6A6"},
    )
    p = box.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INSERT HMI SCREENSHOT HERE")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)
    p2 = box.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(description)
    r2.italic = True
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)

    for style_name in ("List Bullet",):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.5)

    # Cover page
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(1.45))

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("HMI INTERFACE DESCRIPTION")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ORMIN / ORMECO Monitoring System")
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    document.add_paragraph()
    info = document.add_table(rows=6, cols=2)
    info.alignment = 1
    info.autofit = False
    info.columns[0].width = Cm(5.0)
    info.columns[1].width = Cm(8.0)
    cover_rows = [
        ("Customer", "[Customer Name]"),
        ("Project", "[Project / Vessel Name]"),
        ("Document Number", "[Document Number]"),
        ("Revision", "0.1 – Initial Submission"),
        ("Date", "[DD-MMM-YYYY]"),
        ("Prepared by", "[Company / Prepared by]"),
    ]
    for row, values in zip(info.rows, cover_rows):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell, top=100, bottom=100)
            if idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            rr = cell.paragraphs[0].add_run(value)
            rr.bold = idx == 0
            rr.font.name = "Arial"
            rr.font.size = Pt(9.5)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": "B4C6E7"},
                bottom={"val": "single", "sz": "4", "color": "B4C6E7"},
                left={"val": "single", "sz": "4", "color": "B4C6E7"},
                right={"val": "single", "sz": "4", "color": "B4C6E7"},
            )

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(50)
    r = p.add_run("INITIAL CUSTOMER SUBMISSION")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)

    # Main content starts in a section with header/footer.
    main = document.add_section(WD_SECTION.NEW_PAGE)
    main.page_width = Cm(21.0)
    main.page_height = Cm(29.7)
    main.top_margin = Cm(2.1)
    main.bottom_margin = Cm(1.7)
    main.left_margin = Cm(2.0)
    main.right_margin = Cm(2.0)
    set_repeat_header(main, document)

    add_title(document, "1. Purpose")
    add_body(
        document,
        "This document provides a brief description of the Human–Machine Interface (HMI) supplied for the ORMIN / ORMECO monitoring system. The HMI provides live equipment monitoring, historical trends, fuel-consumption analysis and system-status information.",
    )

    add_title(document, "2. System Interface")
    add_body(document, "The HMI is a web-based application connected to the monitoring system through the following interfaces:")
    add_bullets(
        document,
        [
            "Field communication: Modbus TCP",
            "Live-data processing: HMI backend server",
            "Historical-data storage: SQLite database",
            "Operator interface: Standard web browser",
            "Data update: Automatic periodic refresh",
        ],
    )

    architecture = document.add_table(rows=1, cols=4)
    architecture.alignment = 1
    architecture.autofit = False
    labels = ["Sensors / PLC", "Modbus TCP", "HMI Backend", "Web HMI Display"]
    fills = [LIGHT_BLUE, "E2F0D9", "FFF2CC", LIGHT_BLUE]
    for cell, label, fill in zip(architecture.rows[0].cells, labels, fills):
        cell.width = Cm(4.1)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=180, bottom=180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "color": BLUE},
            bottom={"val": "single", "sz": "6", "color": BLUE},
            left={"val": "single", "sz": "6", "color": BLUE},
            right={"val": "single", "sz": "6", "color": BLUE},
        )

    add_title(document, "3. HMI Screens")
    screens = [
        ("Overview", "Summary of the main operating parameters for D/G No. 1 to D/G No. 4."),
        ("Engine", "Detailed operating parameters for the selected engine."),
        ("Engine Status", "Engine operating and status indications."),
        ("Device Status", "Operating condition of monitored equipment and devices."),
        ("P&ID", "Process values and equipment status on the system diagram."),
        ("Pressure Trend", "Historical pressure values compared with engine load."),
        ("Exhaust Temperature Trend", "Cylinder and turbocharger temperatures compared with engine load."),
        ("Fuel Consumption", "D.O. and H.O. inlet flow, return flow and calculated consumption."),
        ("Administration", "Authorized signal, collector and database configuration functions."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(12.0)
    headers = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for cell, label in zip(headers, ("Screen", "Description")):
        set_cell_shading(cell, NAVY)
        rr = cell.paragraphs[0].add_run(label)
        rr.bold = True
        rr.font.name = "Arial"
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor.from_string(WHITE)
    for screen, description in screens:
        cells = table.add_row().cells
        cells[0].text = screen
        cells[1].text = description
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            set_cell_margins(cell, top=65, bottom=65)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "3", "color": "BFBFBF"},
                bottom={"val": "single", "sz": "3", "color": "BFBFBF"},
                left={"val": "single", "sz": "3", "color": "BFBFBF"},
                right={"val": "single", "sz": "3", "color": "BFBFBF"},
            )

    add_title(document, "4. Displayed Information")
    add_body(document, "The HMI displays operating information including:")
    add_bullets(
        document,
        [
            "Engine power, speed and running hours",
            "Fuel-oil, lubricating-oil, cooling-water and intake-air parameters",
            "Cylinder exhaust and turbocharger temperatures",
            "Fuel inlet, return and transfer flow",
            "Equipment running, stopped, trip and fault conditions",
            "Communication status and historical trends",
        ],
    )

    add_title(document, "5. Status and Communication")
    add_body(
        document,
        "The HMI uses values, text, symbols and colours to identify normal, stopped, warning, trip, fault and communication-unavailable conditions. Live information is acquired through Modbus TCP. Historical information is stored locally for trend and fuel-consumption display.",
    )
    add_body(
        document,
        "If communication is interrupted, the HMI provides a communication-status indication. Displayed information must not be considered current until communication is restored.",
    )

    add_title(document, "6. Scope and Confirmation")
    add_body(
        document,
        "The current HMI is intended primarily for monitoring, visualization and operational analysis. Final signal addresses, equipment descriptions, engineering units, status colours, alarm limits and screen arrangements are subject to customer review and project commissioning.",
    )

    # Editable screenshot placeholders.
    shots = document.add_section(WD_SECTION.NEW_PAGE)
    shots.page_width = Cm(21.0)
    shots.page_height = Cm(29.7)
    shots.top_margin = Cm(2.1)
    shots.bottom_margin = Cm(1.7)
    shots.left_margin = Cm(2.0)
    shots.right_margin = Cm(2.0)
    set_repeat_header(shots, document)
    add_title(document, "Appendix A – HMI Screen Captures")
    add_body(document, "Replace each grey placeholder with the applicable approved HMI screenshot before final submission.")

    placeholder_items = screens
    for idx, (screen, description) in enumerate(placeholder_items):
        add_screenshot_placeholder(document, f"A.{idx + 1}  {screen}", description)
        if idx % 2 == 1 and idx != len(placeholder_items) - 1:
            document.add_page_break()

    document.core_properties.title = "HMI Interface Description"
    document.core_properties.subject = "ORMIN / ORMECO HMI System"
    document.core_properties.author = "DRUMS"
    document.core_properties.comments = "Initial customer submission template with editable screenshot placeholders."
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
